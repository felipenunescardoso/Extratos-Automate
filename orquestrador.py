import calendar
import unicodedata
import importlib
from pathlib import Path
from datetime import date, datetime
import re

from openpyxl import load_workbook

from inicial import conectar
from extrator_pdf_generico import extrair_pdf_generico
from extrator_xls_generico import extrair_xls_generico
from extrator_xlsx_generico import extrair_xlsx_generico




def normalizar_chave(texto):
    """
    Normaliza nomes de campos somente para comparação.

    Remove:
        - acentos;
        - cedilha;
        - espaços;
        - hífens;
        - underscores;
        - parênteses;
        - símbolos;
        - diferença entre maiúsculas/minúsculas.

    Exemplos:
        Número Documento -> NUMERODOCUMENTO
        Numero Documento -> NUMERODOCUMENTO
        Descritivo Bancário -> DESCRITIVOBANCARIO
    """

    texto = str(texto or "").strip().upper()

    texto = unicodedata.normalize(
        "NFKD",
        texto
    )

    texto = "".join(
        caractere
        for caractere in texto
        if not unicodedata.combining(
            caractere
        )
    )

    texto = re.sub(
        r"[^A-Z0-9]",
        "",
        texto
    )

    return texto


# ============================================================
# CONSULTAS AO POSTGRESQL
# ============================================================

def buscar_diretorios(conexao):

    with conexao.cursor() as cursor:

        cursor.execute("""
            SELECT
                id,
                diretorio,
                tipo,
                usuario,
                senha,
                finalidade,
                nome_arquivo
            FROM extratos.diretorios
            WHERE ativo = TRUE
            ORDER BY id
        """)

        return cursor.fetchall()

def buscar_arquivos_bancos(conexao):

    with conexao.cursor() as cursor:

        cursor.execute("""
            SELECT
                id,
                nome,
                prefixo,
                extensao,
                abas,
                id_configuracao,
                extrator
            FROM extratos.arquivos_bancos
            WHERE ativo = TRUE
            ORDER BY id
        """)

        return cursor.fetchall()

def buscar_empresa(conexao, nome_empresa):

    with conexao.cursor() as cursor:

        cursor.execute("""
            SELECT
                id_empresa,
                nome_empresa
            FROM extratos.empresas
            WHERE UPPER(nome_empresa) = UPPER(%s)
            LIMIT 1
        """, (nome_empresa,))

        return cursor.fetchone()


def cadastrar_empresa(conexao, nome_empresa):

    with conexao.cursor() as cursor:

        cursor.execute("""
            INSERT INTO extratos.empresas
                (nome_empresa)
            VALUES
                (%s)
            RETURNING
                id_empresa,
                nome_empresa
        """, (nome_empresa,))

        empresa = cursor.fetchone()

    conexao.commit()

    print(
        f"[EMPRESA CADASTRADA] "
        f"{nome_empresa}"
    )

    return empresa


def garantir_empresa(conexao, nome_empresa):

    empresa = buscar_empresa(
        conexao,
        nome_empresa
    )

    if empresa:
        return empresa

    return cadastrar_empresa(
        conexao,
        nome_empresa
    )


def buscar_de_para(conexao, id_arquivo):

    with conexao.cursor() as cursor:

        cursor.execute("""
            SELECT
                nome_campo_arquivo,
                nome_campo_planilha_mae
            FROM extratos.de_para_banco_planilha_mae
            WHERE id_arquivo = %s
        """, (id_arquivo,))

        registros = cursor.fetchall()

    de_para = {}

    for registro in registros:

        nome_campo_arquivo = registro[0]
        nome_campo_planilha = registro[1]

        chave_arquivo = normalizar_chave(
            nome_campo_arquivo
        )

        chave_planilha = normalizar_chave(
            nome_campo_planilha
        )

        de_para[
            chave_arquivo
        ] = chave_planilha

    return de_para


def buscar_campo_origem_por_destino(
    de_para,
    nome_campo_planilha
):
    chave_destino = normalizar_chave(
        nome_campo_planilha
    )

    for (
        campo_origem,
        campo_destino
    ) in de_para.items():

        if campo_destino == chave_destino:
            return campo_origem

    return None


# ============================================================
# CONTROLE DO SERVIÇO / LOG
# ============================================================

ACOES_VALIDAS = {"INSERT", "UPDATE", "DELETE"}


def normalizar_acao(acao):
    acao = str(acao or "").strip().upper()

    if acao not in ACOES_VALIDAS:
        raise ValueError(
            f"Ação inválida em controle_servico: {acao}. "
            f"Use INSERT, UPDATE ou DELETE."
        )

    return acao


def buscar_execucao_pendente(conexao, id_arquivo):
    with conexao.cursor() as cursor:
        cursor.execute("""
            SELECT
                id_execucao,
                data_inicio,
                data_fim,
                acao,
                data_log,
                executado,
                id_arquivo
            FROM extratos.controle_servico
            WHERE id_arquivo = %s
              AND executado = FALSE
            ORDER BY id_execucao
            LIMIT 1
        """, (id_arquivo,))
        registro = cursor.fetchone()

    if not registro:
        return None

    return {
        "id_execucao": registro[0],
        "data_inicio": registro[1],
        "data_fim": registro[2],
        "acao": normalizar_acao(registro[3]),
        "data_log": registro[4],
        "executado": registro[5],
        "id_arquivo": registro[6],
    }


def buscar_ultima_execucao_insert(conexao):
    with conexao.cursor() as cursor:
        cursor.execute("""
            SELECT
                id_execucao,
                data_inicio,
                data_fim
            FROM extratos.controle_servico
            WHERE executado = TRUE
              AND UPPER(acao) = 'INSERT'
            ORDER BY data_fim DESC, id_execucao DESC
            LIMIT 1
        """)
        registro = cursor.fetchone()

    if not registro:
        return None

    return {
        "id_execucao": registro[0],
        "data_inicio": registro[1],
        "data_fim": registro[2],
    }


def criar_execucao(conexao, data_inicio, data_fim, acao="INSERT"):
    acao = normalizar_acao(acao)

    with conexao.cursor() as cursor:
        cursor.execute("""
            INSERT INTO extratos.controle_servico
            (
                data_inicio,
                data_fim,
                acao,
                executado
            )
            VALUES (%s, %s, %s, FALSE)
            RETURNING id_execucao
        """, (
            data_inicio,
            data_fim,
            acao
        ))

        id_execucao = cursor.fetchone()[0]

    conexao.commit()
    return id_execucao


def finalizar_execucao(conexao, id_execucao):
    with conexao.cursor() as cursor:
        cursor.execute("""
            UPDATE extratos.controle_servico
               SET executado = TRUE,
                   data_log = CURRENT_TIMESTAMP
             WHERE id_execucao = %s
        """, (id_execucao,))

    conexao.commit()


def registrar_log(conexao, acao, mensagem):
    acao = normalizar_acao(acao)

    with conexao.cursor() as cursor:
        cursor.execute("""
            INSERT INTO extratos.log
            (
                acao,
                log_execucao
            )
            VALUES (%s, %s)
        """, (
            acao,
            mensagem
        ))

    conexao.commit()


def _como_data(valor):
    if valor is None:
        return None

    if isinstance(valor, datetime):
        return valor.date()

    if isinstance(valor, date):
        return valor

    texto = str(valor).strip()

    if not texto:
        return None

    for formato in (
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%d-%m-%Y"
    ):
        try:
            return datetime.strptime(
                texto,
                formato
            ).date()
        except ValueError:
            pass

    return None


def inicio_fim_mes(data_referencia):
    data_referencia = _como_data(data_referencia)

    if data_referencia is None:
        raise ValueError(
            "Data inválida para calcular o período mensal."
        )

    ultimo_dia = calendar.monthrange(
        data_referencia.year,
        data_referencia.month
    )[1]

    return (
        date(
            data_referencia.year,
            data_referencia.month,
            1
        ),
        date(
            data_referencia.year,
            data_referencia.month,
            ultimo_dia
        )
    )


def filtrar_registros_periodo(
    registros,
    data_inicio,
    data_fim,
    de_para
):
    inicio = _como_data(data_inicio)
    fim = _como_data(data_fim)

    if inicio is None or fim is None:
        raise ValueError(
            "Período de execução inválido."
        )

    if fim < inicio:
        raise ValueError(
            "data_fim não pode ser menor que data_inicio."
        )

    campo_data_origem = buscar_campo_origem_por_destino(
        de_para,
        "Data"
    )

    if campo_data_origem is None:
        raise ValueError(
            "Não existe campo de origem configurado "
            "no DE/PARA para a coluna 'Data' "
            "da Planilha Mãe."
        )

    filtrados = []

    for registro in registros:

        registro_normalizado = {
            normalizar_chave(chave): valor
            for chave, valor in registro.items()
        }

        data_registro = _como_data(
            registro_normalizado.get(
                campo_data_origem
            )
        )

        if (
            data_registro is not None
            and inicio <= data_registro <= fim
        ):
            filtrados.append(registro)

    return filtrados


def excluir_periodo_planilha(
    ws,
    colunas_planilha,
    data_inicio,
    data_fim,
    prefixo_banco=None
):
    chave_data = normalizar_chave(
        "Data"
    )

    if chave_data not in colunas_planilha:
        raise ValueError(
            "Coluna 'Data' não encontrada na Planilha Mãe."
        )

    chave_banco = normalizar_chave(
        "Banco"
    )

    if (
        prefixo_banco is not None
        and chave_banco not in colunas_planilha
    ):
        raise ValueError(
            "Coluna 'Banco' não encontrada na Planilha Mãe."
        )

    inicio = _como_data(data_inicio)
    fim = _como_data(data_fim)

    if inicio is None or fim is None:
        raise ValueError(
            "Período inválido para exclusão."
        )

    coluna_data = colunas_planilha[
        chave_data
    ]

    coluna_banco = (
        colunas_planilha[chave_banco]
        if prefixo_banco is not None
        else None
    )

    prefixo_comparacao = (
        normalizar_chave(prefixo_banco)
        if prefixo_banco is not None
        else None
    )

    quantidade = 0

    for linha in range(
        ws.max_row,
        1,
        -1
    ):
        valor_data = ws.cell(
            row=linha,
            column=coluna_data
        ).value

        data_linha = _como_data(
            valor_data
        )

        if (
            data_linha is None
            or not (
                inicio
                <= data_linha
                <= fim
            )
        ):
            continue

        if prefixo_comparacao is not None:
            valor_banco = ws.cell(
                row=linha,
                column=coluna_banco
            ).value

            banco_linha = normalizar_chave(
                valor_banco
            )

            if banco_linha != prefixo_comparacao:
                continue

        ws.delete_rows(
            linha,
            1
        )

        quantidade += 1

    return quantidade

def descobrir_proximo_periodo(
    conexao,
    diretorios,
    layouts,
    ultima_data_fim=None
):
    limite = _como_data(ultima_data_fim)
    menor_data = None

    for diretorio in diretorios:
        raiz = Path(diretorio[1])

        if not raiz.exists():
            continue

        for pasta_empresa in raiz.iterdir():
            if not pasta_empresa.is_dir():
                continue

            for arquivo in pasta_empresa.rglob("*"):
                if not arquivo.is_file():
                    continue

                layout = identificar_layout(
                    arquivo,
                    layouts
                )

                if layout is None:
                    continue

                de_para = buscar_de_para(
                    conexao,
                    layout["id"]
                )

                if not de_para:
                    continue

                layout_dict = {
                    "id": layout[0],
                    "nome": layout[1],
                    "prefixo": layout[2],
                    "extensao": str(layout[3] or "").strip().lower().replace(".", ""),
                    "abas": layout[4],
                    "id_configuracao": layout[5],
                    "extrator": layout[6]
                }

                nome_extrator = str(
                    layout_dict.get("extrator") or ""
                ).strip()

                try:
                    if nome_extrator in (
                        "extrator_pdf_generico",
                        "extrator_xls_generico",
                        "extrator_xlsx_generico"
                    ):
                        registros = extrair_layout_generico(
                            arquivo,
                            layout_dict
                        )
                    else:
                        funcao_extrator = carregar_extrator(
                            layout_dict
                        )
                        registros = funcao_extrator(
                            arquivo,
                            list(de_para.keys())
                        )
                except (ImportError, ValueError, NotImplementedError):
                    continue

                campo_data_origem = (
                    buscar_campo_origem_por_destino(
                        de_para,
                        "Data"
                    )
                )

                if campo_data_origem is None:
                    continue

                for registro in registros:

                    registro_normalizado = {
                        normalizar_chave(chave): valor
                        for chave, valor in registro.items()
                    }

                    data_registro = _como_data(
                        registro_normalizado.get(
                            campo_data_origem
                        )
                    )

                    if data_registro is None:
                        continue

                    if (
                        limite is not None
                        and data_registro <= limite
                    ):
                        continue

                    if (
                        menor_data is None
                        or data_registro < menor_data
                    ):
                        menor_data = data_registro

    if menor_data is None:
        return None

    return inicio_fim_mes(menor_data)



def extrair_layout_generico(
    arquivo,
    layout
):
    """
    Usado SOMENTE quando:
        - o prefixo não é SSG, ST, BB ou IT;
        - id_arquivo > 4.

    A extensão cadastrada em extratos.arquivos_bancos
    define qual extrator genérico será chamado.
    """

    id_arquivo = int(
        layout["id"]
    )

    if id_arquivo <= 4:
        raise NotImplementedError(
            f"Layout ID {id_arquivo} não é elegível "
            f"para extrator genérico."
        )

    extensao = str(
        layout["extensao"]
        or ""
    ).strip().lower().replace(
        ".",
        ""
    )

    abas = layout.get(
        "abas"
    )

    if extensao == "pdf":

        return extrair_pdf_generico(
            arquivo,
            abas=abas
        )

    if extensao == "xls":

        return extrair_xls_generico(
            arquivo,
            abas=abas
        )

    if extensao == "xlsx":

        return extrair_xlsx_generico(
            arquivo,
            abas=abas
        )

    raise NotImplementedError(
        f"Extensão '{extensao}' não possui "
        f"extrator genérico configurado."
    )


# ============================================================
# IDENTIFICAÇÃO DO ARQUIVO
# ============================================================

def buscar_prefixo_por_id_arquivo(
    layouts,
    id_arquivo
):
    if id_arquivo is None:
        return None

    for layout in layouts:
        if layout[0] == id_arquivo:
            return str(
                layout[2]
            ).strip()

    raise ValueError(
        f"id_arquivo {id_arquivo} não encontrado "
        f"em extratos.arquivos_bancos."
    )


def identificar_layout(arquivo, layouts):

    nome_arquivo = arquivo.name.upper()

    extensao_arquivo = (
        arquivo.suffix
        .lower()
        .replace(".", "")
    )

    for layout in layouts:

        id_arquivo = layout[0]
        nome = layout[1]
        prefixo = layout[2]
        extensao = layout[3]
        abas = layout[4]
        id_configuracao = layout[5]
        extrator = layout[6]

        prefixo_comparacao = (
            str(prefixo or "")
            .strip()
            .upper()
        )

        extensao_comparacao = (
            str(extensao or "")
            .strip()
            .lower()
            .replace(".", "")
        )

        if (
            nome_arquivo.startswith(
                prefixo_comparacao
            )
            and extensao_arquivo
            == extensao_comparacao
        ):

            return {
                "id": id_arquivo,
                "nome": nome,
                "prefixo": prefixo,
                "extensao": extensao_comparacao,
                "abas": abas,
                "id_configuracao": id_configuracao,
                "extrator": extrator
            }

    return None

def localizar_colunas_planilha(ws):

    colunas = {}

    for celula in ws[1]:

        if celula.value is None:
            continue

        nome = str(
            celula.value
        ).strip()

        chave = normalizar_chave(
            nome
        )

        colunas[
            chave
        ] = celula.column

    return colunas


def adicionar_registros_planilha(
    ws,
    colunas_planilha,
    empresa,
    prefixo,
    registros,
    de_para,
    data_inicio=None,
    data_fim=None
):

    # ------------------------------------------------
    # COLUNA DA ENTIDADE
    # ------------------------------------------------
    # Postos usa normalmente "Posto".
    # Distribuidoras pode usar "Empresa" ou "Distribuidora".
    # O primeiro nome existente na Planilha Mãe é utilizado.
    nomes_entidade = (
        "Posto",
        "Empresa",
        "Distribuidora"
    )

    chave_entidade = None

    for nome_entidade in nomes_entidade:
        chave_teste = normalizar_chave(
            nome_entidade
        )

        if chave_teste in colunas_planilha:
            chave_entidade = chave_teste
            break

    if chave_entidade is None:
        raise ValueError(
            "Não foi encontrada coluna para a empresa "
            "na Planilha Mãe. Esperado: Posto, Empresa "
            "ou Distribuidora."
        )

    chave_banco = normalizar_chave(
        "Banco"
    )

    if chave_banco not in colunas_planilha:
        raise ValueError(
            "Coluna 'Banco' não encontrada "
            "na Planilha Mãe."
        )

    quantidade = 0

    campo_valor_origem = buscar_campo_origem_por_destino(
        de_para,
        "Valor"
    )

    if campo_valor_origem is None:
        raise ValueError(
            "Não existe campo de origem configurado "
            "no DE/PARA para a coluna 'Valor' "
            "da Planilha Mãe."
        )

    if (
        data_inicio is not None
        and data_fim is not None
    ):
        registros = filtrar_registros_periodo(
            registros,
            data_inicio,
            data_fim,
            de_para
        )

    for registro in registros:

        registro_normalizado = {
            normalizar_chave(chave): valor
            for chave, valor in registro.items()
        }

        if registro_normalizado.get(
            campo_valor_origem
        ) is None:
            continue

        linha_excel = ws.max_row + 1

        ws.cell(
            row=linha_excel,
            column=colunas_planilha[
                chave_entidade
            ],
            value=empresa
        )

        ws.cell(
            row=linha_excel,
            column=colunas_planilha[
                chave_banco
            ],
            value=prefixo
        )

        for (
            campo_arquivo,
            campo_planilha
        ) in de_para.items():

            if (
                campo_arquivo
                not in registro_normalizado
            ):
                continue

            if (
                campo_planilha
                not in colunas_planilha
            ):
                print(
                    f"[AVISO] Coluna configurada "
                    f"no DE/PARA não encontrada "
                    f"na Planilha Mãe: "
                    f"{campo_planilha}"
                )
                continue

            valor = registro_normalizado[
                campo_arquivo
            ]

            ws.cell(
                row=linha_excel,
                column=colunas_planilha[
                    campo_planilha
                ],
                value=valor
            )

        quantidade += 1

    return quantidade

def carregar_extrator(layout):
    nome_modulo = str(layout.get("extrator") or "").strip()

    if not nome_modulo:
        raise ValueError(
            f"Layout ID {layout['id']} não possui extrator configurado "
            "em extratos.arquivos_bancos."
        )

    modulo = importlib.import_module(nome_modulo)

    funcoes = [
        getattr(modulo, nome)
        for nome in dir(modulo)
        if nome.startswith("extrair_")
        and callable(getattr(modulo, nome))
        and getattr(modulo, nome).__module__ == modulo.__name__
    ]

    if len(funcoes) != 1:
        raise ValueError(
            f"Extrator '{nome_modulo}' deve possuir exatamente uma "
            "função pública iniciada por 'extrair_'. "
            f"Encontradas: {len(funcoes)}."
        )

    return funcoes[0]


def processar_arquivo(
    conexao,
    arquivo,
    empresa,
    layout,
    ws,
    colunas_planilha,
    data_inicio,
    data_fim
):

    print(
        f"    [ARQUIVO] "
        f"{arquivo.name}"
    )

    # --------------------------------------------
    # BUSCA CONFIGURAÇÃO DE/PARA
    # --------------------------------------------

    de_para = buscar_de_para(
        conexao,
        layout["id"]
    )

    if not de_para:

        print(
            f"    [PULADO] "
            f"Sem DE/PARA configurado."
        )

        return {
            "inseridos": 0,
            "processado": False
        }

    nome_extrator = str(
        layout.get("extrator") or ""
    ).strip()

    if nome_extrator in (
        "extrator_pdf_generico",
        "extrator_xls_generico",
        "extrator_xlsx_generico"
    ):
        registros = extrair_layout_generico(
            arquivo,
            layout
        )
    else:
        funcao_extrator = carregar_extrator(
            layout
        )

        try:
            registros = funcao_extrator(
                arquivo,
                list(de_para.keys())
            )

        except ValueError as erro:
            mensagem_erro = str(erro)

            prefixo_erro = (
                str(
                    layout.get(
                        "prefixo",
                        ""
                    )
                )
                .strip()
                .upper()
            )

            # REGRA GENÉRICA:
            # qualquer banco pode possuir arquivos/contas sem novos
            # lançamentos para o fluxo de conta corrente.
            # Essa condição não deve interromper a fila.
            if (
                "Já existem lançamentos na data."
                in mensagem_erro
            ):
                print(
                    f"    [PULADO] "
                    f"{arquivo.name}: {mensagem_erro}"
                )

                return {
                    "inseridos": 0,
                    "processado": False
                }

            # Arquivos dentro da pasta RENDIMENTO que não pertencem
            # aos layouts bancários suportados não devem derrubar
            # toda a execução do mês. Apenas pula o arquivo.
            if (
                prefixo_erro == "REND"
                and (
                    "Layout do Banco Rendimento não reconhecido"
                    in mensagem_erro
                    or "Layout Internet Banking do Rendimento"
                    in mensagem_erro
                )
            ):
                print(
                    f"    [PULADO] "
                    f"{arquivo.name}: {mensagem_erro}"
                )

                return {
                    "inseridos": 0,
                    "processado": False
                }

            # SAFRA:
            # PDFs que não são extrato de conta corrente
            # ou cujo layout ainda não é suportado não devem
            # interromper o restante da execução.
            if (
                prefixo_erro == "SAF"
                and (
                    "Layout do Banco Safra não reconhecido"
                    in mensagem_erro
                    or "Não foi possível identificar o ano "
                    "do período no extrato Safra"
                    in mensagem_erro
                )
            ):
                print(
                    f"    [PULADO] "
                    f"{arquivo.name}: {mensagem_erro}"
                )

                return {
                    "inseridos": 0,
                    "processado": False
                }

            raise

    # --------------------------------------------
    # ESCREVE NA PLANILHA
    # --------------------------------------------

    quantidade = (
        adicionar_registros_planilha(
            ws=ws,
            colunas_planilha=colunas_planilha,
            empresa=empresa,
            prefixo=layout["prefixo"],
            registros=registros,
            de_para=de_para,
            data_inicio=data_inicio,
            data_fim=data_fim
        )
    )

    print(
        f"    [OK] "
        f"{quantidade} lançamentos."
    )

    return {
        "inseridos": quantidade,
        "processado": quantidade > 0
    }


# ============================================================
# CONFIGURAÇÃO / ESTRUTURA DE DIRETÓRIOS
# ============================================================

def buscar_layout_por_id(
    layouts,
    id_arquivo
):
    if id_arquivo is None:
        return None

    for layout in layouts:
        if int(layout[0]) == int(id_arquivo):
            return {
                "id": layout[0],
                "nome": layout[1],
                "prefixo": layout[2],
                "extensao": str(
                    layout[3] or ""
                ).strip().lower().replace(".", ""),
                "abas": layout[4],
                "id_configuracao": layout[5],
                "extrator": layout[6]
            }

    raise ValueError(
        f"id_arquivo {id_arquivo} não encontrado "
        f"em extratos.arquivos_bancos."
    )


def planilha_mae_do_diretorio(
    diretorio
):
    raiz = Path(diretorio[1])

    nome_arquivo = str(
        diretorio[6] or ""
    ).strip()

    if not nome_arquivo:
        raise ValueError(
            f"Diretório ID {diretorio[0]} não possui "
            "nome_arquivo configurado."
        )

    caminho = raiz / nome_arquivo

    # Compatibilidade caso o cadastro tenha sido feito
    # sem a extensão visível no nome.
    if not caminho.suffix:
        caminho = caminho.with_suffix(
            ".xlsx"
        )

    return caminho


def meses_do_periodo(
    data_inicio,
    data_fim
):
    inicio = _como_data(data_inicio)
    fim = _como_data(data_fim)

    atual = date(
        inicio.year,
        inicio.month,
        1
    )

    limite = date(
        fim.year,
        fim.month,
        1
    )

    meses = []

    while atual <= limite:
        meses.append(
            (
                atual.year,
                atual.month
            )
        )

        if atual.month == 12:
            atual = date(
                atual.year + 1,
                1,
                1
            )
        else:
            atual = date(
                atual.year,
                atual.month + 1,
                1
            )

    return meses


def localizar_pasta_mes(
    raiz,
    ano,
    mes
):
    # A estrutura atual usa:
    # 01 - JANEIRO
    # 02 - FEVEREIRO
    # ...
    padrao = re.compile(
        rf"^\s*{mes:02d}\s*[-_]"
    )

    for pasta in raiz.iterdir():
        if (
            pasta.is_dir()
            and padrao.search(
                pasta.name.upper()
            )
        ):
            return pasta

    return None


def localizar_pasta_banco(
    pasta_mes,
    layout
):
    prefixo = normalizar_chave(
        layout["prefixo"]
    )

    nome = normalizar_chave(
        layout["nome"]
    )

    nome_base = normalizar_chave(
        str(layout["nome"] or "").split(" - ")[0]
    )

    for pasta in pasta_mes.iterdir():
        if not pasta.is_dir():
            continue

        chave = normalizar_chave(
            pasta.name
        )

        if chave in (
            prefixo,
            nome,
            nome_base
        ):
            return pasta

    return None


def arquivo_compativel_layout(
    arquivo,
    layout
):
    extensao = (
        arquivo.suffix
        .lower()
        .replace(".", "")
    )

    return (
        arquivo.is_file()
        and extensao
        == layout["extensao"]
    )



def arquivo_eh_aplicacao(nome_arquivo):
    """
    Identifica arquivos de conta/aplicação independentemente do banco.

    Exemplos cobertos:
        APLIC ...
        APL AUT ...
        APLIC AUTO ...
        ... APLIC ...
        ... APL AUT ...
        APLICAÇÃO / APLICACAO ...

    Esses arquivos não pertencem ao fluxo de extrato de conta corrente
    da Planilha Mãe e devem ser ignorados antes do extrator.
    """

    nome = str(nome_arquivo or "").strip().upper()

    nome = unicodedata.normalize(
        "NFKD",
        nome
    )

    nome = "".join(
        caractere
        for caractere in nome
        if not unicodedata.combining(caractere)
    )

    nome = re.sub(
        r"[^A-Z0-9]+",
        " ",
        nome
    ).strip()

    padroes = (
        r"\bAPLIC\b",
        r"\bAPLICACAO\b",
        r"\bAPL\s+AUT\b",
        r"\bAPLIC\s+AUT\b",
        r"\bAPLIC\s+AUTO\b",
        r"\bINVESTIMENTO\b",
        r"\bGARANTIDA\b",
    )

    return any(
        re.search(
            padrao,
            nome
        )
        for padrao in padroes
    )


def normalizar_empresa_distribuidora(
    nome_arquivo,
    pasta_empresa=None
):
    nome_pasta = str(pasta_empresa or "").strip()
    if nome_pasta:
        chave_pasta = normalizar_chave(nome_pasta)
        if chave_pasta == "PRIME":
            return "PRIME"
        if chave_pasta == "EURO":
            return "EURO"

    nome = str(nome_arquivo or "").strip()
    nome = re.sub(r"\.pdf$", "", nome, flags=re.IGNORECASE).strip()
    nome = re.sub(r"\s*\(\d+\)\s*$", "", nome).strip()
    nome = re.sub(r"\s+BB\s*-\s*$", "", nome, flags=re.IGNORECASE).strip()

    # FEATBANK - março:
    # AG1 CC3812-1_ECOMAR IMPORTADORA LTDA -> ECOMAR
    # EXTRATO_AG1 CC0003816-4_VENTURE T I E P SA -> VENTURE
    match_feat = re.match(
        r"^(?:EXTRATO_)?AG1\s+CC\d+(?:-\d+)?_(.+)$",
        nome,
        flags=re.IGNORECASE
    )

    if match_feat:
        nome_feat = match_feat.group(1).strip()

        normalizacoes_feat = (
            ("ECOMAR", "ECOMAR"),
            ("GRAN BAHIA", "GRAN"),
            ("GREEN FICUS", "GREEN"),
            ("PARAISO", "PARAISO"),
            ("REFINARIA DE PETROLEO DE PILAR", "PILAR"),
            ("EZ PAY", "EZPAY"),
            ("LOGFIT", "LOGFIT"),
            ("NORTE SUL", "NORTE SUL"),
            ("XFS", "XFS"),
            ("BCAG", "BCAG"),
            ("NEW POWER", "NEW"),
            ("CHRONOS", "CHRONOS"),
            ("ATHENA", "ATHENA"),
            ("TIF", "TIF"),
            ("T LIQ", "TLIQ"),
            ("TLIQ", "TLIQ"),
            ("EURO", "EURO"),
            ("MARANGONI", "MARANGONI"),
            ("VENTURE", "VENTURE"),
            ("TERMINAL DE CAMPOS", "TECAM"),
            ("SHERIDAN", "SHERIDAN"),
        )

        nome_feat_upper = nome_feat.upper()

        for inicio_nome, nome_normalizado in normalizacoes_feat:
            if nome_feat_upper.startswith(inicio_nome):
                return nome_normalizado

    nome = re.sub(
        r"^\s*\d{3,}(?:(?:-[0-9X]+)|X)?\s+",
        "",
        nome,
        flags=re.IGNORECASE
    ).strip()
    nome = re.sub(
        r"\s+\d{3,}(?:(?:-[0-9X]+)|X)?$",
        "",
        nome,
        flags=re.IGNORECASE
    ).strip()

    # Reaplica a limpeza depois da remoção da conta.
    nome = re.sub(
        r"\s+BB\s*-\s*$",
        "",
        nome,
        flags=re.IGNORECASE
    ).strip()

    nome_upper = nome.upper()
    correcoes_exatas = {
        "FOCOUS": "FOCUS",
    }
    if nome_upper in correcoes_exatas:
        return correcoes_exatas[nome_upper]

    if re.match(r"^76", nome, flags=re.IGNORECASE):
        return "76 OIL"
    if re.search(r"\bEURO\s+G(?:\b|ESTAO)", nome, flags=re.IGNORECASE):
        return "EURO"
    if re.search(r"\bPRIME(?:\s+FACTOR)?\b", nome, flags=re.IGNORECASE):
        return "PRIME"
    if re.match(r"^EASY", nome, flags=re.IGNORECASE):
        return "EASYPETRO"

    return nome


# ============================================================
# PROCESSAMENTO DO DIRETÓRIO
# ============================================================


def processar_diretorio(
    conexao,
    diretorio,
    layouts,
    acao,
    data_inicio,
    data_fim,
    id_arquivo=None
):
    acao = normalizar_acao(acao)

    id_configuracao = diretorio[0]
    raiz = Path(diretorio[1])
    finalidade = str(
        diretorio[5] or ""
    ).strip().upper()

    if not raiz.exists():
        print("[ERRO] Diretório não encontrado:")
        print(raiz)
        return {
            "inseridos": 0,
            "excluidos": 0,
            "arquivos_processados": 0
        }

    layout_especifico = buscar_layout_por_id(
        layouts,
        id_arquivo
    )

    # Uma ordem só pode ser executada no diretório
    # correspondente ao id_configuracao do layout.
    if (
        layout_especifico is not None
        and int(
            layout_especifico["id_configuracao"]
        ) != int(id_configuracao)
    ):
        return {
            "inseridos": 0,
            "excluidos": 0,
            "arquivos_processados": 0
        }

    print()
    print("=" * 70)
    print(
        f"DIRETÓRIO: {raiz}"
    )
    print(
        f"FINALIDADE: {finalidade}"
    )
    print(
        f"EXECUÇÃO: {acao} | "
        f"{data_inicio:%d/%m/%Y} até "
        f"{data_fim:%d/%m/%Y}"
    )
    print("=" * 70)

    planilha_mae = planilha_mae_do_diretorio(
        diretorio
    )

    if not planilha_mae.exists():
        print(
            "[ERRO] Planilha Mãe não encontrada:"
        )
        print(planilha_mae)
        return {
            "inseridos": 0,
            "excluidos": 0,
            "arquivos_processados": 0
        }

    print(
        f"[PLANILHA MÃE] "
        f"{planilha_mae}"
    )

    workbook = load_workbook(
        planilha_mae
    )

    ws = workbook.active

    colunas_planilha = (
        localizar_colunas_planilha(
            ws
        )
    )

    total_inseridos = 0
    total_excluidos = 0
    total_arquivos_processados = 0

    prefixo_banco = (
        buscar_prefixo_por_id_arquivo(
            layouts,
            id_arquivo
        )
    )

    if acao in (
        "UPDATE",
        "DELETE"
    ):
        total_excluidos = (
            excluir_periodo_planilha(
                ws=ws,
                colunas_planilha=colunas_planilha,
                data_inicio=data_inicio,
                data_fim=data_fim,
                prefixo_banco=prefixo_banco
            )
        )

        print(
            f"[PERÍODO REMOVIDO] "
            f"{total_excluidos} linhas"
            + (
                f" do banco {prefixo_banco}."
                if prefixo_banco
                else "."
            )
        )

    if acao == "DELETE":
        workbook.save(
            planilha_mae
        )

        print(
            f"[PLANILHA SALVA] "
            f"{planilha_mae}"
        )

        return {
            "inseridos": 0,
            "excluidos": total_excluidos,
            "arquivos_processados": 0
        }

    # ========================================================
    # DISTRIBUIDORAS
    # raiz / 01 - JANEIRO / BB / arquivos
    # ========================================================
    if finalidade == "DISTRIBUIDORAS":

        if layout_especifico is None:
            raise ValueError(
                "Distribuidoras exige id_arquivo "
                "na ordem de controle_servico."
            )

        for ano, mes in meses_do_periodo(
            data_inicio,
            data_fim
        ):
            inicio_mes = date(
                ano,
                mes,
                1
            )

            ultimo_dia_mes = calendar.monthrange(
                ano,
                mes
            )[1]

            fim_mes = date(
                ano,
                mes,
                ultimo_dia_mes
            )

            pasta_mes = localizar_pasta_mes(
                raiz,
                ano,
                mes
            )

            if pasta_mes is None:
                print(
                    f"[MÊS NÃO ENCONTRADO] "
                    f"{mes:02d}/{ano}"
                )
                continue

            print()
            print(
                f"[MÊS] {pasta_mes.name}"
            )

            pasta_banco = localizar_pasta_banco(
                pasta_mes,
                layout_especifico
            )

            if pasta_banco is None:
                print(
                    f"[BANCO NÃO ENCONTRADO] "
                    f"{layout_especifico['prefixo']} "
                    f"em {pasta_mes.name}"
                )
                continue

            print(
                f"[BANCO] "
                f"{pasta_banco.name}"
            )

            for arquivo in pasta_banco.rglob(
                "*"
            ):
                if not arquivo_compativel_layout(
                    arquivo,
                    layout_especifico
                ):
                    continue

                # REGRA GENÉRICA:
                # contas/arquivos de aplicação podem existir em
                # qualquer banco e não entram no fluxo da Planilha Mãe.
                if arquivo_eh_aplicacao(
                    arquivo.stem
                ):
                    print(
                        f"    [PULADO] "
                        f"{arquivo.name}: "
                        f"arquivo de aplicação"
                    )
                    continue

                # SAFRA:
                # notas de negociação também não pertencem
                # ao extrato de conta corrente.
                # Ignora antes de cadastrar empresa/processar.
                if (
                    str(
                        layout_especifico.get(
                            "prefixo",
                            ""
                        )
                    ).strip().upper() == "SAF"
                ):
                    nome_arquivo_upper = (
                        arquivo.stem.strip().upper()
                    )

                    if (
                        nome_arquivo_upper.startswith(
                            "NOTA DE NEGOCIA"
                        )
                        or nome_arquivo_upper.startswith(
                            "NOTAS DE NEGOCIA"
                        )
                    ):
                        print(
                            f"    [PULADO] "
                            f"{arquivo.name}: "
                            f"arquivo de negociação"
                        )
                        continue

                # Nas distribuidoras o arquivo fica dentro
                # da pasta do banco. O nome do arquivo
                # identifica a empresa.
                pasta_empresa = None

                try:
                    relativo_banco = arquivo.relative_to(
                        pasta_banco
                    )

                    if len(relativo_banco.parts) > 1:
                        pasta_empresa = relativo_banco.parts[0]

                except ValueError:
                    pasta_empresa = None

                nome_empresa = (
                    normalizar_empresa_distribuidora(
                        arquivo.stem,
                        pasta_empresa=pasta_empresa
                    )
                )

                garantir_empresa(
                    conexao,
                    nome_empresa
                )

                resultado_arquivo = (
                    processar_arquivo(
                        conexao=conexao,
                        arquivo=arquivo,
                        empresa=nome_empresa,
                        layout=layout_especifico,
                        ws=ws,
                        colunas_planilha=colunas_planilha,
                        data_inicio=inicio_mes,
                        data_fim=fim_mes
                    )
                )

                total_inseridos += (
                    resultado_arquivo[
                        "inseridos"
                    ]
                )

                if resultado_arquivo[
                    "processado"
                ]:
                    total_arquivos_processados += 1

    # ========================================================
    # POSTOS / ESTRUTURA ANTIGA
    # raiz / empresa / arquivos
    # ========================================================
    else:

        layouts_configuracao = [
            layout
            for layout in layouts
            if int(layout[5])
            == int(id_configuracao)
        ]

        for pasta_empresa in raiz.iterdir():
            if not pasta_empresa.is_dir():
                continue

            nome_empresa = (
                pasta_empresa.name.strip()
            )

            garantir_empresa(
                conexao,
                nome_empresa
            )

            print()
            print(
                f"[EMPRESA] "
                f"{nome_empresa}"
            )

            for arquivo in pasta_empresa.rglob(
                "*"
            ):
                if not arquivo.is_file():
                    continue

                layout = identificar_layout(
                    arquivo,
                    layouts_configuracao
                )

                if layout is None:
                    continue

                if (
                    id_arquivo is not None
                    and layout["id"]
                    != id_arquivo
                ):
                    continue

                resultado_arquivo = (
                    processar_arquivo(
                        conexao=conexao,
                        arquivo=arquivo,
                        empresa=nome_empresa,
                        layout=layout,
                        ws=ws,
                        colunas_planilha=colunas_planilha,
                        data_inicio=data_inicio,
                        data_fim=data_fim
                    )
                )

                total_inseridos += (
                    resultado_arquivo[
                        "inseridos"
                    ]
                )

                if resultado_arquivo[
                    "processado"
                ]:
                    total_arquivos_processados += 1

    workbook.save(
        planilha_mae
    )

    print()
    print(
        f"[PLANILHA SALVA] "
        f"{planilha_mae}"
    )

    return {
        "inseridos": total_inseridos,
        "excluidos": total_excluidos,
        "arquivos_processados": total_arquivos_processados
    }

def executar_execucao(
    conexao,
    diretorios,
    layouts,
    execucao
):
    id_execucao = execucao["id_execucao"]

    data_inicio = _como_data(
        execucao["data_inicio"]
    )

    data_fim = _como_data(
        execucao["data_fim"]
    )

    acao = normalizar_acao(
        execucao["acao"]
    )

    id_arquivo = execucao.get(
        "id_arquivo"
    )

    if (
        data_inicio is None
        or data_fim is None
    ):
        raise ValueError(
            f"Execução {id_execucao} possui período inválido."
        )

    if data_fim < data_inicio:
        raise ValueError(
            f"Execução {id_execucao}: "
            "data_fim menor que data_inicio."
        )

    print()
    print("=" * 70)
    print(
        f"EXECUÇÃO {id_execucao} | {acao}"
    )
    print(
        f"PERÍODO: "
        f"{data_inicio:%d/%m/%Y} até "
        f"{data_fim:%d/%m/%Y}"
    )

    if id_arquivo is not None:
        print(
            f"ARQUIVO CONFIGURADO: ID {id_arquivo}"
        )

    print("=" * 70)

    registrar_log(
        conexao,
        acao,
        (
            f"Início da execução {id_execucao}. "
            f"Período {data_inicio:%d/%m/%Y} "
            f"a {data_fim:%d/%m/%Y}."
        )
    )

    total_inseridos = 0
    total_excluidos = 0
    total_arquivos_processados = 0

    for diretorio in diretorios:

        resultado = processar_diretorio(
            conexao=conexao,
            diretorio=diretorio,
            layouts=layouts,
            acao=acao,
            data_inicio=data_inicio,
            data_fim=data_fim,
            id_arquivo=id_arquivo
        )

        total_inseridos += resultado["inseridos"]
        total_excluidos += resultado["excluidos"]
        total_arquivos_processados += resultado[
            "arquivos_processados"
        ]

    if (
        id_arquivo is not None
        and acao in ("INSERT", "UPDATE")
        and total_arquivos_processados == 0
    ):
        raise ValueError(
            f"Execução {id_execucao}: nenhum lançamento "
            f"foi inserido para o id_arquivo {id_arquivo} "
            f"no período informado. "
            f"A execução permanecerá pendente."
        )

    finalizar_execucao(
        conexao,
        id_execucao
    )

    registrar_log(
        conexao,
        acao,
        (
            f"Execução {id_execucao} concluída com sucesso. "
            f"Período {data_inicio:%d/%m/%Y} "
            f"a {data_fim:%d/%m/%Y}. "
            f"Inseridos: {total_inseridos}. "
            f"Excluídos: {total_excluidos}. "
            f"Arquivos processados: {total_arquivos_processados}."
        )
    )

    print()
    print("=" * 70)
    print("PROCESSAMENTO FINALIZADO")
    print(f"Execução: {id_execucao}")
    print(f"Ação: {acao}")
    print(
        f"Total inserido: {total_inseridos}"
    )
    print(
        f"Total excluído: {total_excluidos}"
    )
    print(
        f"Arquivos processados: "
        f"{total_arquivos_processados}"
    )
    print("=" * 70)


def main():

    conexao = None
    execucao = None

    try:

        with conectar() as conexao:

            diretorios = buscar_diretorios(
                conexao
            )

            layouts = buscar_arquivos_bancos(
                conexao
            )

            if not diretorios:
                print(
                    "Nenhum diretório ativo cadastrado."
                )
                return

            if not layouts:
                print(
                    "Nenhum arquivo de banco ativo cadastrado."
                )
                return

            encontrou_pendencia = False

            # Primeiro percorre somente os arquivos/layouts ativos.
            # Para cada um, executa apenas serviços com executado = FALSE.
            # Se não houver FALSE para aquele arquivo, ele é pulado.
            for layout in layouts:
                id_arquivo = layout[0]
                nome_layout = layout[1]

                execucao = buscar_execucao_pendente(
                    conexao,
                    id_arquivo
                )

                if execucao is None:
                    print(
                        f"[PULADO] ID {id_arquivo} - {nome_layout}: "
                        "sem serviço pendente (executado = FALSE)."
                    )
                    continue

                encontrou_pendencia = True

                # Vida normal: processa todas as pendências FALSE
                # desse arquivo ativo, em ordem de id_execucao.
                while execucao is not None:
                    executar_execucao(
                        conexao=conexao,
                        diretorios=diretorios,
                        layouts=layouts,
                        execucao=execucao
                    )

                    execucao = buscar_execucao_pendente(
                        conexao,
                        id_arquivo
                    )

            if not encontrou_pendencia:
                print()
                print(
                    "[SEM PENDÊNCIAS] "
                    "Nenhum arquivo ativo possui serviço com "
                    "executado = FALSE."
                )
                return

            print()
            print(
                "[FILA CONCLUÍDA] "
                "Nenhuma execução pendente."
            )

    except Exception as erro:

        print()
        print("[ERRO NO PROCESSAMENTO]")
        print(erro)

        if (
            conexao is not None
            and execucao is not None
        ):
            try:
                acao_execucao = normalizar_acao(
                    execucao["acao"]
                )

                registrar_log(
                    conexao,
                    acao_execucao,
                    (
                        f"Erro na execução "
                        f"{execucao.get('id_execucao')}: "
                        f"{erro}"
                    )
                )
            except Exception:
                pass



if __name__ == "__main__":
    main()