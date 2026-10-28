from openpyxl import load_workbook


def _interpretar_abas(abas):
    if abas is None:
        return None

    if isinstance(abas, (list, tuple, set)):
        nomes = [
            str(valor).strip()
            for valor in abas
            if str(valor).strip()
        ]
        return nomes or None

    texto = str(abas).strip()

    if not texto:
        return None

    texto = texto.replace(",", ";")

    nomes = [
        valor.strip()
        for valor in texto.split(";")
        if valor.strip()
    ]

    return nomes or None


def _normalizar_cabecalhos(valores):
    resultado = []
    usados = {}

    for indice, valor in enumerate(valores, start=1):

        nome = str(valor or "").strip()

        if not nome:
            nome = f"COLUNA_{indice}"

        quantidade = usados.get(nome, 0) + 1
        usados[nome] = quantidade

        if quantidade > 1:
            nome = f"{nome}_{quantidade}"

        resultado.append(nome)

    return resultado


def extrair_xlsx_generico(caminho, abas=None):
    """
    Extrator genérico XLSX.
    A primeira linha não vazia de cada aba é usada como cabeçalho.
    """

    workbook = load_workbook(
        caminho,
        data_only=True,
        read_only=True
    )

    abas_configuradas = _interpretar_abas(abas)

    if abas_configuradas:
        nomes_abas = [
            nome
            for nome in abas_configuradas
            if nome in workbook.sheetnames
        ]
    else:
        nomes_abas = list(workbook.sheetnames)

    registros = []

    for nome_aba in nomes_abas:

        ws = workbook[nome_aba]
        linhas_uteis = []

        for linha in ws.iter_rows(values_only=True):

            if any(
                valor is not None
                and str(valor).strip() != ""
                for valor in linha
            ):
                linhas_uteis.append(list(linha))

        if len(linhas_uteis) < 2:
            continue

        cabecalhos = _normalizar_cabecalhos(linhas_uteis[0])

        for linha in linhas_uteis[1:]:

            valores = list(linha)

            if len(valores) < len(cabecalhos):
                valores.extend(
                    [None] * (len(cabecalhos) - len(valores))
                )

            if len(valores) > len(cabecalhos):
                valores = valores[:len(cabecalhos)]

            registro = {
                cabecalhos[i]: valor
                for i, valor in enumerate(valores)
            }

            registro["_Aba"] = nome_aba
            registros.append(registro)

    workbook.close()

    return registros