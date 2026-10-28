from openpyxl import load_workbook


#CAMINHO_PLANILHA = r"D:\AutomateBancos\EXTRATOS MENSAIS\EXTRATOS MENSAIS\Planilha Mae Postos.xlsx"

CAMINHO_PLANILHA = r"D:\AutomateBancos\Dist Extratos Mensais\Extratos Mensais\Planilha Mae Dist.xlsx"


def apagar_dados():

    bancos_digitados = input(
        "Quais bancos deseja apagar? "
        "Ex: BR, BB, SAF\n"
        "Pressione ENTER para apagar TODOS: "
    ).strip()

    workbook = load_workbook(
        CAMINHO_PLANILHA
    )

    worksheet = workbook.active

    ultima_linha = worksheet.max_row
    ultima_coluna = worksheet.max_column

    if ultima_linha <= 1:

        print(
            "[OK] A planilha já está sem dados."
        )

        workbook.close()
        return

    # ========================================================
    # SEM BANCO INFORMADO -> APAGA TUDO
    # ========================================================

    if not bancos_digitados:

        quantidade = ultima_linha - 1

        # Remove todo o bloco de dados de uma vez.
        worksheet.delete_rows(
            2,
            quantidade
        )

        workbook.save(
            CAMINHO_PLANILHA
        )

        workbook.close()

        print(
            f"[OK] {quantidade} linhas apagadas."
        )

        print(
            "[OK] Cabeçalho da linha 1 preservado."
        )

        return

    # ========================================================
    # APAGA SOMENTE OS BANCOS INFORMADOS
    # ========================================================

    bancos = {
        banco.strip().upper()
        for banco in bancos_digitados.split(",")
        if banco.strip()
    }

    coluna_banco = None

    for celula in worksheet[1]:

        if (
            str(celula.value or "")
            .strip()
            .upper()
            == "BANCO"
        ):

            coluna_banco = celula.column
            break

    if coluna_banco is None:

        workbook.close()

        raise ValueError(
            "Coluna 'Banco' não encontrada "
            "na Planilha Mãe."
        )

    # --------------------------------------------------------
    # Em vez de apagar linha por linha com delete_rows(),
    # lê os dados uma única vez, mantém somente os bancos
    # desejados e depois recria o bloco de dados.
    #
    # Isso evita milhares de deslocamentos internos no Excel.
    # --------------------------------------------------------

    linhas_manter = []
    quantidade = 0

    for valores in worksheet.iter_rows(
        min_row=2,
        max_row=ultima_linha,
        max_col=ultima_coluna,
        values_only=True
    ):

        banco_linha = (
            str(
                valores[coluna_banco - 1]
                or ""
            )
            .strip()
            .upper()
        )

        if banco_linha in bancos:
            quantidade += 1
            continue

        linhas_manter.append(
            list(valores)
        )

    if quantidade == 0:

        workbook.close()

        print(
            f"[OK] Nenhuma linha encontrada "
            f"para: {', '.join(sorted(bancos))}"
        )

        return

    # Remove TODAS as linhas de dados de uma vez.
    worksheet.delete_rows(
        2,
        ultima_linha - 1
    )

    # Reinsere somente os bancos que devem permanecer.
    for valores in linhas_manter:
        worksheet.append(
            valores
        )

    workbook.save(
        CAMINHO_PLANILHA
    )

    workbook.close()

    print(
        f"[OK] Bancos apagados: "
        f"{', '.join(sorted(bancos))}"
    )

    print(
        f"[OK] {quantidade} linhas apagadas."
    )

    print(
        f"[OK] {len(linhas_manter)} linhas preservadas."
    )

    print(
        "[OK] Demais bancos preservados."
    )


if __name__ == "__main__":
    apagar_dados()
